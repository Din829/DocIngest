"""
Unit tests for xlsx drawing-layer arrow attribution (2026-07-11):

Callout labels (text boxes) are tied to their targets by drawn arrows.
The label TEXT already survives via 〔図形: …〕, but its anchor row is
where the box sits — measured on a real 扶養控除 form, one label sat 18
rows from where its arrow pointed, so "下記のとおりです" landed next to
the wrong table block. The collector now pairs each directed arrow to
the label box it starts at (absolute-EMU geometry) and records what it
points at (anchor from/to cells — Excel's own EMU→cell truth), and the
parser appends "→ 対象: <target>" to the rendered label.

Geometry validated on the real form first (18/18 arrows, scratchpad
prototype); these tests lock the contract with hand-built OOXML:
  * flip combinations decide which bounding-box corners the line joins
  * tailEnd vs headEnd decides the pointed end
  * cxnSp connectors take the same path as sp lines
  * double-headed / headless lines, custGeom, oneCellAnchor → skipped
  * pairing: flush hit / far start / two-box ambiguity → refused
  * arrow endpoint on another label box → {"kind": "shape"} target
  * group-nested shapes: text still collected, excluded from pairing
  * _xlsx_arrow_target_label: left-walk + row-scan, never inflates dims

Pure in-memory zips + openpyxl — no parsing pipeline, no network.

Run:
    python tests/unit/test_xlsx_arrow_attribution.py
"""

from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT / "src"))

import openpyxl

from docingest.parsers.docling_parser import (
    _collect_xlsx_image_anchors,
    _xlsx_arrow_target_label,
)

XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
M = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


# ---------------------------------------------------------------------------
# OOXML builders
# ---------------------------------------------------------------------------

def _sp_box(text, x, y, cx, cy, sp_id="1"):
    """Top-level text box with an absolute-EMU xfrm."""
    return (
        f'<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="{sp_id}" name="b"/>'
        f"<xdr:cNvSpPr/></xdr:nvSpPr>"
        f"<xdr:spPr><a:xfrm><a:off x=\"{x}\" y=\"{y}\"/>"
        f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom></xdr:spPr>'
        f"<xdr:txBody><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p>"
        f"</xdr:txBody></xdr:sp>"
    )


def _sp_arrow(x, y, cx, cy, geom="line", head="none", tail="triangle",
              flip=""):
    """Text-less line shape; flip like 'flipH="1" flipV="1"'."""
    return (
        f'<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="2" name="a"/>'
        f"<xdr:cNvSpPr/></xdr:nvSpPr>"
        f"<xdr:spPr><a:xfrm {flip}><a:off x=\"{x}\" y=\"{y}\"/>"
        f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>'
        f'<a:ln><a:headEnd type="{head}"/><a:tailEnd type="{tail}"/></a:ln>'
        f"</xdr:spPr></xdr:sp>"
    )


def _cxn_arrow(x, y, cx, cy, flip="", geom="straightConnector1",
               st_id=None, end_id=None):
    """Connector shape carrying a tail arrow, optionally snapped to shapes."""
    cxns = ""
    if st_id is not None:
        cxns += f'<a:stCxn id="{st_id}" idx="0"/>'
    if end_id is not None:
        cxns += f'<a:endCxn id="{end_id}" idx="0"/>'
    return (
        f'<xdr:cxnSp><xdr:nvCxnSpPr><xdr:cNvPr id="99" name="c"/>'
        f"<xdr:cNvCxnSpPr>{cxns}</xdr:cNvCxnSpPr></xdr:nvCxnSpPr>"
        f"<xdr:spPr><a:xfrm {flip}><a:off x=\"{x}\" y=\"{y}\"/>"
        f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>'
        f'<a:ln><a:tailEnd type="triangle"/></a:ln>'
        f"</xdr:spPr></xdr:cxnSp>"
    )


def _anchor(frm, to, inner):
    """twoCellAnchor; frm/to are 0-based (row, col) as Excel writes them."""
    (fr, fc), (tr, tc) = frm, to
    return (
        f"<xdr:twoCellAnchor><xdr:from><xdr:col>{fc}</xdr:col>"
        f"<xdr:colOff>0</xdr:colOff><xdr:row>{fr}</xdr:row>"
        f"<xdr:rowOff>0</xdr:rowOff></xdr:from>"
        f"<xdr:to><xdr:col>{tc}</xdr:col><xdr:colOff>0</xdr:colOff>"
        f"<xdr:row>{tr}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>"
        f"{inner}<xdr:clientData/></xdr:twoCellAnchor>"
    )


def _one_cell_anchor(frm, inner):
    (fr, fc) = frm
    return (
        f"<xdr:oneCellAnchor><xdr:from><xdr:col>{fc}</xdr:col>"
        f"<xdr:colOff>0</xdr:colOff><xdr:row>{fr}</xdr:row>"
        f"<xdr:rowOff>0</xdr:rowOff></xdr:from>"
        f'<xdr:ext cx="100000" cy="100000"/>'
        f"{inner}<xdr:clientData/></xdr:oneCellAnchor>"
    )


def _write_xlsx(path, drawing_body):
    drawing = f'<xdr:wsDr xmlns:xdr="{XDR}" xmlns:a="{A}">{drawing_body}</xdr:wsDr>'
    parts = {
        "[Content_Types].xml":
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            "</Types>",
        "_rels/.rels":
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{R}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml":
            f'<?xml version="1.0"?><workbook xmlns="{M}" xmlns:r="{R}">'
            '<sheets><sheet name="設計" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels":
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{R}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml":
            f'<?xml version="1.0"?><worksheet xmlns="{M}" xmlns:r="{R}">'
            '<sheetData/><drawing r:id="rId1"/></worksheet>',
        "xl/worksheets/_rels/sheet1.xml.rels":
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{R}/drawing" Target="../drawings/drawing1.xml"/></Relationships>',
        "xl/drawings/drawing1.xml": drawing,
    }
    with zipfile.ZipFile(path, "w") as z:
        for name, content in parts.items():
            z.writestr(name, content)


def _collect(drawing_body):
    fd, p = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        _write_xlsx(p, drawing_body)
        return _collect_xlsx_image_anchors(Path(p)).get("設計", [])
    finally:
        os.unlink(p)


def _entry(entries, text):
    for e in entries:
        if e.get("text") == text:
            return e
    raise AssertionError(f"entry {text!r} not found in {entries}")


# Shared fixture: label box at EMU (1_000_000, 1_000_000)-(3_000_000, 2_000_000),
# anchored at 0-based (3, 0). Arrow starts flush under the box.
BOX = _anchor((3, 0), (5, 3), _sp_box("注意A", 1_000_000, 1_000_000, 2_000_000, 1_000_000))


# ---------------------------------------------------------------------------
# Direction: flips and arrow-head side
# ---------------------------------------------------------------------------

def test_no_flip_points_to_anchor_to():
    # Arrow bbox starts inside the box, no flip → line runs top-left →
    # bottom-right, tail arrow → points at the anchor's ``to`` cell.
    arrow = _anchor((4, 1), (8, 3),
                    _sp_arrow(2_000_000, 1_900_000, 500_000, 1_500_000))
    entries = _collect(BOX + arrow)
    e = _entry(entries, "注意A")
    assert e["targets"] == [{"kind": "cell", "row": 9, "col": 4}], e
    print("  ✓ no flip: target = anchor to-cell")


def test_flip_v_swaps_rows():
    # flipV → line runs bottom-left → top-right: the arrow END is on the
    # FROM row. Box must sit at the line's start (bbox bottom edge).
    arrow = _anchor((4, 1), (8, 3),
                    _sp_arrow(2_000_000, 2_100_000, 500_000, 1_500_000,
                              flip='flipV="1"'))
    entries = _collect(BOX + arrow)
    e = _entry(entries, "注意A")
    # EMU start = (x, y+cy) = (2_000_000, 3_600_000)… that's 1.6M below the
    # box → NOT paired. Instead anchor the bbox so its bottom touches the
    # box: y such that y+cy ≈ 2_000_000.
    assert "targets" not in e, e
    arrow = _anchor((4, 1), (8, 3),
                    _sp_arrow(2_000_000, 550_000, 500_000, 1_500_000,
                              flip='flipV="1"'))
    entries = _collect(BOX + arrow)
    e = _entry(entries, "注意A")
    assert e["targets"] == [{"kind": "cell", "row": 5, "col": 4}], e
    print("  ✓ flipV: rows swap (end row = anchor from-row); start follows bbox")


def test_head_arrow_reverses_direction():
    # Arrow head on headEnd → pointed end is the line START; the box side
    # is the line END. No flip: line start = bbox top-left = the pointed
    # cell (anchor from), so the box must touch the bbox BOTTOM-RIGHT.
    arrow = _anchor((8, 3), (4, 1),
                    _sp_arrow(500_000, 2_000_000, 1_400_000, 550_000,
                              head="triangle", tail="none"))
    # bbox bottom-right = (1_900_000, 2_550_000) — within TOL of the box
    # bottom edge (y2=2_000_000)? dy=550_000 > TOL. Put it flush instead:
    arrow = _anchor((8, 3), (4, 1),
                    _sp_arrow(500_000, 900_000, 1_400_000, 550_000,
                              head="triangle", tail="none"))
    # bottom-right = (1_900_000, 1_450_000): inside the box → paired;
    # pointed end = anchor FROM cell (0-based (8,3) → 1-based (9,4)).
    entries = _collect(BOX + arrow)
    e = _entry(entries, "注意A")
    assert e["targets"] == [{"kind": "cell", "row": 9, "col": 4}], e
    print("  ✓ headEnd-only arrow: direction reversed, points at from-cell")


def test_cxnsp_connector_same_path():
    arrow = _anchor((4, 1), (8, 3),
                    _cxn_arrow(2_000_000, 1_900_000, 500_000, 1_500_000))
    entries = _collect(BOX + arrow)
    e = _entry(entries, "注意A")
    assert e["targets"] == [{"kind": "cell", "row": 9, "col": 4}], e
    print("  ✓ cxnSp connector attributed like an sp line")


# ---------------------------------------------------------------------------
# Skips: not-an-arrow shapes
# ---------------------------------------------------------------------------

def test_double_headed_and_headless_skipped():
    both = _anchor((4, 1), (8, 3),
                   _sp_arrow(2_000_000, 1_900_000, 500_000, 1_500_000,
                             head="triangle", tail="triangle"))
    none_ = _anchor((4, 1), (8, 3),
                    _sp_arrow(2_000_000, 1_900_000, 500_000, 1_500_000,
                              tail="none"))
    entries = _collect(BOX + both + none_)
    assert "targets" not in _entry(entries, "注意A"), entries
    print("  ✓ double-headed and headless lines attach nothing")


def test_custgeom_and_onecellanchor_skipped():
    freeform = _anchor((4, 1), (8, 3),
                       _sp_arrow(2_000_000, 1_900_000, 500_000, 1_500_000,
                                 geom="custGeom"))
    one_cell = _one_cell_anchor((4, 1),
                                _sp_arrow(2_000_000, 1_900_000, 500_000,
                                          1_500_000))
    entries = _collect(BOX + freeform + one_cell)
    assert "targets" not in _entry(entries, "注意A"), entries
    print("  ✓ custGeom and oneCellAnchor arrows skipped")


# ---------------------------------------------------------------------------
# Pairing: tolerance and ambiguity
# ---------------------------------------------------------------------------

def test_far_start_not_paired():
    arrow = _anchor((10, 5), (12, 6),
                    _sp_arrow(6_000_000, 6_000_000, 500_000, 500_000))
    entries = _collect(BOX + arrow)
    assert "targets" not in _entry(entries, "注意A"), entries
    print("  ✓ arrow starting far from every box attaches nothing")


def test_equidistant_boxes_refused():
    # Second box to the right; arrow start exactly between the two edges.
    box_b = _anchor((3, 5), (5, 8),
                    _sp_box("注意B", 3_100_000, 1_000_000, 900_000, 1_000_000))
    # Start x=3_050_000: 50_000 from A's right edge AND from B's left edge.
    arrow = _anchor((4, 4), (8, 5),
                    _sp_arrow(3_050_000, 1_500_000, 500_000, 1_500_000))
    entries = _collect(BOX + box_b + arrow)
    assert "targets" not in _entry(entries, "注意A"), entries
    assert "targets" not in _entry(entries, "注意B"), entries
    print("  ✓ two nearly-equidistant boxes → attribution refused")


def test_box_to_box_target():
    box_b = _anchor((10, 8), (12, 10),
                    _sp_box("手順2へ", 8_000_000, 8_000_000, 1_000_000, 1_000_000))
    # Starts flush under 注意A, ends inside 手順2へ.
    arrow = _anchor((4, 1), (11, 9),
                    _sp_arrow(2_000_000, 1_900_000, 6_500_000, 6_600_000))
    entries = _collect(BOX + box_b + arrow)
    e = _entry(entries, "注意A")
    assert e["targets"] == [{"kind": "shape", "text": "手順2へ"}], e
    print("  ✓ arrow ending on another label box → shape target")


def test_multiple_arrows_one_box():
    a1 = _anchor((4, 1), (8, 3),
                 _sp_arrow(2_000_000, 1_900_000, 500_000, 1_500_000))
    a2 = _anchor((4, 1), (10, 5),
                 _sp_arrow(1_200_000, 1_900_000, 900_000, 2_500_000))
    entries = _collect(BOX + a1 + a2)
    e = _entry(entries, "注意A")
    kinds = sorted((t["row"], t["col"]) for t in e["targets"])
    assert kinds == [(9, 4), (11, 6)], e
    print("  ✓ one box owning two arrows lists both targets")


# ---------------------------------------------------------------------------
# Explicit connections (stCxn/endCxn) — authoritative over geometry.
# Modelled on a real flowchart (入社手続き): bentConnector whose flip-derived
# direction was WRONG vs. the rendered truth; the attachments are right.
# ---------------------------------------------------------------------------

def test_explicit_connection_overrides_geometry():
    # Boxes far apart; connector bbox placed so that GEOMETRY alone would
    # pair its flip-derived start with box A — but st/endCxn say B → A.
    box_a = _anchor((3, 0), (5, 3),
                    _sp_box("ラベルA", 1_000_000, 1_000_000, 2_000_000, 1_000_000,
                            sp_id="10"))
    box_b = _anchor((10, 8), (12, 10),
                    _sp_box("ラベルB", 8_000_000, 8_000_000, 1_000_000, 1_000_000,
                            sp_id="11"))
    arrow = _anchor((4, 1), (11, 9),
                    _cxn_arrow(2_000_000, 1_900_000, 6_500_000, 6_600_000,
                               st_id="11", end_id="10"))
    entries = _collect(box_a + box_b + arrow)
    e = _entry(entries, "ラベルB")
    assert e["targets"] == [{"kind": "shape", "text": "ラベルA"}], entries
    assert "targets" not in _entry(entries, "ラベルA"), entries
    print("  ✓ st/endCxn attachments override flip geometry")


def test_bent_connector_half_attached():
    # Real-world case (入社手続き bent1): bent connector, one end snapped
    # to box B, the other dangling near box A. flipV puts the endpoints on
    # the bl/tr diagonal (geometric fact, kept); the flip-derived DIRECTION
    # would call bl the path start — attachment-based orientation must win:
    # tr is the corner nearer B (attached), so the free/owner end is bl.
    box_a = _anchor((3, 0), (5, 3),
                    _sp_box("ラベルA", 1_000_000, 8_600_000, 2_000_000, 1_000_000,
                            sp_id="10"))
    box_b = _anchor((10, 8), (12, 10),
                    _sp_box("ラベルB", 8_000_000, 1_500_000, 1_000_000, 1_000_000,
                            sp_id="11"))
    arrow = _anchor((4, 1), (11, 9),
                    _cxn_arrow(2_000_000, 1_900_000, 6_500_000, 6_600_000,
                               geom="bentConnector2", flip='flipV="1"',
                               end_id="11"))
    entries = _collect(box_a + box_b + arrow)
    e = _entry(entries, "ラベルA")
    assert e["targets"] == [{"kind": "shape", "text": "ラベルB"}], entries
    print("  ✓ half-attached bentConnector: owner from free corner, flips ignored")


def test_bent_connector_unattached_refused():
    arrow = _anchor((4, 1), (8, 3),
                    _cxn_arrow(2_000_000, 1_900_000, 500_000, 1_500_000,
                               geom="bentConnector2"))
    entries = _collect(BOX + arrow)
    assert "targets" not in _entry(entries, "注意A"), entries
    print("  ✓ unattached bentConnector refused (direction unknowable)")


def test_attachment_to_textless_shape_falls_back():
    # Connector snapped to a decorative (text-less) shape: no label to
    # name, so the pointed end falls back to geometry/cell.
    deco = _anchor((10, 8), (12, 10),
                   '<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="20" name="d"/>'
                   "<xdr:cNvSpPr/></xdr:nvSpPr>"
                   '<xdr:spPr><a:xfrm><a:off x="8000000" y="8000000"/>'
                   '<a:ext cx="1000000" cy="1000000"/></a:xfrm>'
                   '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                   "</xdr:spPr></xdr:sp>")
    arrow = _anchor((4, 1), (11, 9),
                    _cxn_arrow(2_000_000, 1_900_000, 6_500_000, 6_600_000,
                               end_id="20"))
    entries = _collect(BOX + deco + arrow)
    e = _entry(entries, "注意A")
    # Owner via geometry (start flush under 注意A); target: pointed corner
    # is inside the deco shape → no label box there → cell fallback.
    assert e["targets"] == [{"kind": "cell", "row": 12, "col": 10}], entries
    print("  ✓ attachment to a text-less shape → cell fallback for the target")


# ---------------------------------------------------------------------------
# Groups: text collected, geometry excluded
# ---------------------------------------------------------------------------

def test_grouped_shapes_text_kept_geometry_excluded():
    grouped = _anchor((4, 1), (8, 3),
        "<xdr:grpSp><xdr:nvGrpSpPr>"
        '<xdr:cNvPr id="9" name="g"/><xdr:cNvGrpSpPr/>'
        "</xdr:nvGrpSpPr><xdr:grpSpPr/>"
        # Child box whose (child-space!) rect happens to overlap 注意A's
        # arrow start — must NOT steal the pairing.
        + _sp_box("組内ラベル", 1_900_000, 1_800_000, 500_000, 500_000)
        # Child arrow — child-space coords, must not be collected.
        + _sp_arrow(2_000_000, 1_900_000, 500_000, 1_500_000)
        + "</xdr:grpSp>")
    arrow = _anchor((4, 1), (8, 3),
                    _sp_arrow(2_000_000, 1_900_000, 500_000, 1_500_000))
    entries = _collect(BOX + grouped + arrow)
    # Grouped label text still collected (existing behaviour).
    grp = _entry(entries, "組内ラベル")
    assert "targets" not in grp, grp
    # The top-level arrow pairs to 注意A, not the (closer in fake child
    # coords) grouped box.
    e = _entry(entries, "注意A")
    assert e["targets"] == [{"kind": "cell", "row": 9, "col": 4}], e
    print("  ✓ grouped shapes: text kept, excluded from pairing")


# ---------------------------------------------------------------------------
# Target label lookup
# ---------------------------------------------------------------------------

def test_target_label_walks_left_and_scans_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B9"] = "A 源泉控除\n対象配偶者"
    assert _xlsx_arrow_target_label(ws, 9, 4) == "A 源泉控除 対象配偶者"
    ws2 = wb.create_sheet()
    ws2["C7"] = "Employee No."
    assert _xlsx_arrow_target_label(ws2, 8, 4) == "Employee No."
    print("  ✓ label lookup: left-walk on merged rows + nearby-row scan")


def test_target_label_never_inflates_dimensions():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    before = (ws.max_row, ws.max_column)
    assert _xlsx_arrow_target_label(ws, 50, 30) is None
    assert (ws.max_row, ws.max_column) == before, (before, ws.max_row, ws.max_column)
    print("  ✓ out-of-range lookup returns None without creating cells")


if __name__ == "__main__":
    for fn in [
        test_no_flip_points_to_anchor_to,
        test_flip_v_swaps_rows,
        test_head_arrow_reverses_direction,
        test_cxnsp_connector_same_path,
        test_double_headed_and_headless_skipped,
        test_custgeom_and_onecellanchor_skipped,
        test_far_start_not_paired,
        test_equidistant_boxes_refused,
        test_box_to_box_target,
        test_multiple_arrows_one_box,
        test_explicit_connection_overrides_geometry,
        test_bent_connector_half_attached,
        test_bent_connector_unattached_refused,
        test_attachment_to_textless_shape_falls_back,
        test_grouped_shapes_text_kept_geometry_excluded,
        test_target_label_walks_left_and_scans_rows,
        test_target_label_never_inflates_dimensions,
    ]:
        fn()
    print("\nALL xlsx arrow attribution TESTS PASSED")
